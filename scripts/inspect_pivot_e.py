"""Inspect pivot column E on Параметры угля."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
XLSM = ROOT / "tmp_decrypted.xlsm"


def main() -> None:
    z = zipfile.ZipFile(XLSM)
    for name in sorted(z.namelist()):
        if name.startswith("xl/pivotTables/pivotTable") and name.endswith(".xml"):
            xml = z.read(name).decode("utf-8")
            loc = re.search(r'location ref="([^"]+)"', xml)
            title = re.search(r'name="([^"]+)"', xml)
            print(f"\n=== {name} name={title.group(1) if title else '?'} loc={loc.group(1) if loc else '?'} ===")
            for df in re.finditer(
                r'<dataField name="([^"]+)" fld="(\d+)"([^/]*)/>',
                xml,
            ):
                extra = df.group(3)
                print(f"  dataField[{df.group(2)}]: {df.group(1)} {extra.strip()}")

    cache_xml = z.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
    fields = re.findall(r'cacheField name="([^"]+)"', cache_xml)
    print("\n=== CACHE FIELDS (index:name) ===")
    for i, name in enumerate(fields):
        print(f"  [{i}] {name}")

    xml2 = z.read("xl/pivotTables/pivotTable2.xml").decode("utf-8")
    for tag in ("rowFields", "colFields"):
        m = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", xml2, re.S)
        if m:
            print(f"\nPIVOT2 {tag}: {m.group(1).strip()}")
    for m in re.finditer(r'<cacheField[^>]*formula="([^"]+)"[^/]*/>', cache_xml):
        print(f"\nCALCULATED FIELD: {m.group(0)}")

    z.close()

    wb = openpyxl.load_workbook(XLSM, data_only=False, keep_vba=True)
    params = wb.worksheets[0]
    detail = wb.worksheets[1]

    print("\n=== PARAMS A3:E6 (formulas/values) ===")
    for r in range(3, 7):
        row = []
        for c in range(1, 6):
            cell = params.cell(r, c)
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                row.append(f"{get_column_letter(c)}{r}={v[:80]}")
            else:
                row.append(f"{get_column_letter(c)}{r}={v!r}")
        print("  ", " | ".join(row))

    wbv = openpyxl.load_workbook(XLSM, data_only=True, keep_vba=True)
    params_v = wbv.worksheets[0]
    print("\n=== PARAMS computed C5,D5,E5 ===")
    for addr in ["C5", "D5", "E5", "C11", "D11", "E11"]:
        print(f"  {addr} = {params_v[addr].value}")
        if params_v[addr].value and params_v["C" + addr[1:]].value:
            c = params_v["C" + addr[1:]].value
            d = params_v["D" + addr[1:]].value
            e = params_v[addr].value
            if c:
                print(f"    D/C = {d/c:.6f}, E = {e}")

    # find columns on detail sheet matching pivot fields 16, 43, 44
    print("\n=== DETAIL sheet headers for pivot source cols ===")
    # source C3:AT27 => col C=3 is first data col
    # cache field index maps to columns starting from C in source
    for idx in [0, 1, 12, 16, 40, 41, 42, 43, 44]:
        col = 3 + idx  # C is index 0? need verify - worksheetSource ref C3:AT27
        # cache fields order should match columns in range C3:AT27 left to right
        if col <= detail.max_column:
            h = detail.cell(3, col).value
            f4 = detail.cell(4, col).value
            print(f"  cache[{idx}] -> col {get_column_letter(col)}: header={h!r}")
            if isinstance(f4, str) and f4.startswith("="):
                print(f"    row4 formula: {f4[:150]}")
            else:
                print(f"    row4 value: {f4!r}")

    # scan detail row3 for coeff/volume keywords
    print("\n=== DETAIL row3 all headers with col letters ===")
    for c in range(1, 50):
        h = detail.cell(3, c).value
        if h:
            print(f"  {get_column_letter(c)}: {h!r}")
    for c in range(40, 48):
        h = detail.cell(3, c).value
        f = detail.cell(4, c).value
        print(f"  {get_column_letter(c)}: {h!r} | row4={str(f)[:100] if f else None}")

    wb.close()
    wbv.close()


if __name__ == "__main__":
    main()
