"""Inspect Уголь_эластика_2026_5.xlsm structure for elasticity scaling."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as col_letter

ROOT = Path(__file__).resolve().parents[1]
XLSM = ROOT / "tmp_decrypted.xlsm"


def main() -> None:
    wb = openpyxl.load_workbook(XLSM, data_only=False, keep_vba=True)
    wbv = openpyxl.load_workbook(XLSM, data_only=True, keep_vba=True)

    sheets = {s.title: s for s in wb.worksheets}
    names = list(sheets)
    print("SHEETS:")
    for i, name in enumerate(names):
        s = sheets[name]
        pivots = len(getattr(s, "_pivots", []) or [])
        print(f"  [{i}] {name!r} rows={s.max_row} cols={s.max_column} pivots={pivots}")

    detail_name = names[1]
    agg_name = names[3]
    params_name = names[0]
    ref_name = names[5]
    coeff_name = names[4]

    detail = sheets[detail_name]
    agg = sheets[agg_name]
    params = sheets[params_name]
    ref = sheets[ref_name]

    print("\n=== 1. БОЛЬШИЕ ОБОРОТЫ (W) ===")
    print(f"Лист: {detail_name!r}, колонка W (23), заголовок row3:")
    print(f"  {detail.cell(3, 23).value}")
    total_w = 0.0
    for r in range(4, detail.max_row + 1):
        w = detail.cell(r, 23).value
        if isinstance(w, (int, float)):
            total_w += w
            if r <= 8:
                print(
                    f"  row{r}: груз={detail.cell(r, 1).value!r}, "
                    f"холдинг={detail.cell(r, 15).value!r}, W={w:,.0f}"
                )
    print(f"  SUM(W4:W{detail.max_row}) = {total_w:,.0f} руб")

    print("\n=== 2. ЭКОНОМИКА: SUMIF (детальный лист) ===")
    print(f"Лист: {detail_name!r}")
    for addr, label in [
        ("AJ4", "выручка/цена FOB по грузу"),
        ("AG4", "перегрузка по холдингу O"),
        ("AF4", "тариф РЖД (масштаб по вагонам Z+AA+AB / X)"),
        ("AE4", "сумма тарифных компонент AC+AD"),
    ]:
        print(f"  {addr} ({label}):")
        print(f"    {detail[addr].value}")

    print("\n=== 2b. ЭКОНОМИКА: среднее по холдингу (упрощённый лист) ===")
    print(f"Лист: {agg_name!r}")
    print(f"  AJ4 было SUMIF, стало фикс: {agg['AJ4'].value}")
    print(f"  detail AJ4: {detail['AJ4'].value}")
    print(f"  ref D9 (источник для agg): {ref['D9'].value}")

    print("\n=== 3. МАРЖИНАЛЬНОСТЬ И k НА СРЕЗ ===")
    print(f"Лист: {detail_name!r}, row4:")
    print(f"  AO4 маржинальность база: {detail['AO4'].value}")
    print(f"  AZ4 k при тарифе AZ3={detail['AZ3'].value}: {str(detail['AZ4'].value)[:200]}")
    print(f"  IS2 сценарный шаг тарифа (=ROUND(IU2,2)): формула {detail['IS2'].value}")
    print(f"  IU2 тариф с Параметры угля: {detail['IU2'].value}")
    print(f"  IS4 = HLOOKUP(IS2, AZ3:IR27, ...): {str(detail['IS4'].value)[:200]}")
    print(f"  IT4 итоговый k с enterprise load: {str(detail['IT4'].value)[:200]}")
    print(f"  AU4 enterprise load: {detail['AU4'].value}")
    print(f"Таблица эластичности: лист {coeff_name!r}, VLOOKUP -> A:B")

    print("\n=== 4. PIVOT (перегруппировка 24 срезов) ===")
    for p in params._pivots:
        print(f"  pivot {p.name!r} at {p.location}")
    z = zipfile.ZipFile(XLSM)
    xml = z.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
    m = re.search(r'worksheetSource ref="([^"]+)" sheet="([^"]+)"', xml)
    if m:
        print(f"  источник pivot1: лист {m.group(2)!r}, диапазон {m.group(1)}")
    z.close()

    print("\n=== 5. ПРОПОРЦИИ НА Параметры угля ===")
    print(f"Лист: {params_name!r}")
    print("  Блок объёмов по холдингам (rows 28-36):")
    for r in range(28, 37):
        b = params.cell(r, 2).value
        c = params.cell(r, 3).value
        d = params.cell(r, 4).value
        formula_d = params.cell(r, 4).value
        if isinstance(formula_d, str) and formula_d.startswith("="):
            d_show = formula_d
        else:
            d_show = d
        print(f"    row{r}: B={b!r}, C={c}, D={d_show}")

    print("\n  Сценарные переключатели (тариф -> IU2 на маршрутах):")
    for addr in ["G11", "J11", "I14", "J14", "I17", "J17"]:
        v = params[addr].value
        print(f"    {addr}: {v}")

    print("\n  Пример пропорции D34:")
    print(f"    B29={params['B29'].value}, D29={params['D29'].value}")
    print(f"    B34={params['B34'].value}, D34 formula={params['D34'].value}")

    wb.close()
    wbv.close()


if __name__ == "__main__":
    main()
