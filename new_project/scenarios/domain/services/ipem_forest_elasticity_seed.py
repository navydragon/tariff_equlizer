"""Seed ElasticitySet 2026 (forest) from IPEM 'Технический лист'."""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from django.contrib.auth import get_user_model
from django.db import transaction

from core.models import CargoGroup
from scenarios.domain.services.base_elasticity_seed import ELASTICITY_SET_NAME
from scenarios.models import (
    ElasticityRule,
    ElasticityRulePoint,
    ElasticitySet,
    Scenario,
)

User = get_user_model()

TECH_SHEET_NAME = "Технический лист"
FOREST_RULE_NAME = "IPEM: Лес"
FOREST_CARGO_GROUP_CODE = 6
FOREST_RULE_POSITION = 32


@dataclass(frozen=True)
class IpemForestElasticitySeedResult:
    elasticity_set_id: int
    rules_upserted: int
    points_upserted: int
    attached_to_scenario: bool


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[4]


def _default_workbook_path() -> Path:
    return _repo_root() / "data" / "ipem" / "Лес_эластика.xlsx"


def _cell_has_value(value) -> bool:
    return value is not None and str(value).strip() != ""


def _load_points_from_block(
    worksheet,
    *,
    marginality_col_0based: int,
    coefficient_col_0based: int,
) -> list[tuple[Decimal, Decimal]]:
    by_marginality: dict[Decimal, Decimal] = {}
    for row in range(2, worksheet.max_row + 1):
        marginality = worksheet.cell(row, marginality_col_0based + 1).value
        coefficient = worksheet.cell(row, coefficient_col_0based + 1).value
        if not _cell_has_value(marginality) or not _cell_has_value(coefficient):
            continue
        if isinstance(coefficient, str) and coefficient.startswith("="):
            continue
        key = Decimal(str(marginality)).quantize(Decimal("0.0001"))
        by_marginality[key] = Decimal(str(coefficient)).quantize(
            Decimal("0.0001"),
        )
    return sorted(by_marginality.items(), key=lambda item: item[0])


def _resolve_cargo_group_by_code(code: int) -> CargoGroup | None:
    return CargoGroup.objects.filter(code=code).first()


def _resolve_elasticity_set(owner: User) -> ElasticitySet:
    existing = (
        ElasticitySet.objects.filter(name=ELASTICITY_SET_NAME)
        .order_by("id")
        .first()
    )
    if existing is not None:
        return existing
    return ElasticitySet.objects.create(author=owner, name=ELASTICITY_SET_NAME)


def _replace_rule_points(
    rule: ElasticityRule,
    points: list[tuple[Decimal, Decimal]],
) -> int:
    ElasticityRulePoint.objects.filter(rule=rule).delete()
    if not points:
        return 0
    ElasticityRulePoint.objects.bulk_create(
        [
            ElasticityRulePoint(rule=rule, marginality=m, coefficient=c)
            for m, c in points
        ],
    )
    return len(points)


@transaction.atomic
def seed_ipem_forest_elasticity_for_scenario(
    scenario: Scenario,
    *,
    author: User | None = None,
    attach: bool = True,
    xlsx_path: Path | None = None,
) -> IpemForestElasticitySeedResult:
    """
    Создаёт/обновляет правило эластичности «IPEM: Лес» по листу
    «Технический лист» (одна кривая на cargo_group=6).
    """
    owner = author or scenario.author
    if owner is None:
        raise ValueError("author is required to seed elasticity set")

    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to seed IPEM forest elasticity",
        ) from exc

    workbook_path = xlsx_path or _default_workbook_path()
    if not workbook_path.exists():
        return IpemForestElasticitySeedResult(
            elasticity_set_id=0,
            rules_upserted=0,
            points_upserted=0,
            attached_to_scenario=False,
        )

    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    worksheet = workbook[TECH_SHEET_NAME]

    points = _load_points_from_block(
        worksheet,
        marginality_col_0based=0,
        coefficient_col_0based=1,
    )

    elasticity_set = _resolve_elasticity_set(owner)

    ElasticityRule.objects.filter(
        elasticity_set=elasticity_set,
        name=FOREST_RULE_NAME,
    ).delete()

    forest_group = _resolve_cargo_group_by_code(FOREST_CARGO_GROUP_CODE)
    rule = ElasticityRule.objects.create(
        elasticity_set=elasticity_set,
        name=FOREST_RULE_NAME,
        position=FOREST_RULE_POSITION,
        cargo_group=forest_group,
    )
    points_upserted = _replace_rule_points(rule, points)

    attached = False
    if attach and not scenario.elasticity_set_id:
        scenario.elasticity_set = elasticity_set
        scenario.save(update_fields=["elasticity_set"])
        attached = True

    return IpemForestElasticitySeedResult(
        elasticity_set_id=elasticity_set.id,
        rules_upserted=1,
        points_upserted=points_upserted,
        attached_to_scenario=attached,
    )

